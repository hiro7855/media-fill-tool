"""产出带红黄标记的复核 Excel;并能读回人工修改后的表(供回填)。

结构:
  第1行:说明(合并)
  第2行:表头
  第3行起:人员记录(媒介资料 + 若单人单元含行程则一并填)
  空行 + 「待认领行程」小标题 + 行程记录(姓名列标黄,待人工填出行人)
"""
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))

from fields import flat_columns  # noqa: E402

RED = PatternFill("solid", fgColor="FFC7CE")      # 格式错误
YELLOW = PatternFill("solid", fgColor="FFEB9C")    # AI 不确定 / 待人工
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
MARKER_FILL = PatternFill("solid", fgColor="FCE4D6")

EXTRA_COLS = ["_来源", "_备注"]
SHEET_NAME = "复核"
HEADER_ROW = 2
DATA_START = 3
TRIP_MARKER = "▼ 待认领行程:下面每行是一段行程,请在「姓名」列填上出行人(与上方某人对应)"


def _write_row(ws, i, rec, cols):
    row = rec["row"]
    errors = rec.get("errors", {})
    uncertain = rec.get("uncertain", set())
    is_trip = rec.get("kind") == "trip"
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(i, j, row.get(c, ""))
        if c in errors:
            cell.fill = RED
            cell.comment = Comment(errors[c][1], "校验")
        elif c in uncertain:
            cell.fill = YELLOW
            if is_trip and c == "姓名":
                cell.comment = Comment("请填写出行人姓名", "待认领")
            else:
                cell.comment = Comment("AI 不确定,请核对", "AI")
    ws.cell(i, len(cols) + 1, rec.get("source", ""))
    ws.cell(i, len(cols) + 2, rec.get("notes", ""))


def write_review(records, out_path):
    cols = flat_columns()
    all_cols = cols + EXTRA_COLS
    persons = [r for r in records if r.get("kind") != "trip"]
    trips = [r for r in records if r.get("kind") == "trip"]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    tip = ("说明:红色=格式可能有误,黄色=AI 不确定或待人工填。请核对并直接修改单元格,"
           "「待认领行程」区需在「姓名」列填出行人。改完保存本文件,再运行「回填」。")
    ws.cell(1, 1, tip)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(all_cols))
    ws.cell(1, 1).font = Font(bold=True, color="9C0006")

    for j, c in enumerate(all_cols, start=1):
        cell = ws.cell(HEADER_ROW, j, c)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    i = DATA_START
    for rec in persons:
        _write_row(ws, i, rec, cols)
        i += 1

    if trips:
        i += 1  # 留一空行
        ws.cell(i, 1, TRIP_MARKER)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=len(all_cols))
        ws.cell(i, 1).font = Font(bold=True, color="833C00")
        ws.cell(i, 1).fill = MARKER_FILL
        i += 1
        for rec in trips:
            _write_row(ws, i, rec, cols)
            i += 1

    for j, c in enumerate(all_cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(26, len(c) + 8))
    ws.freeze_panes = f"A{DATA_START}"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def read_review(path):
    """读回复核表,返回 list of {列名:值}。跳过说明行、空行、「待认领」小标题行。"""
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    headers = [ws.cell(HEADER_ROW, j).value for j in range(1, ws.max_column + 1)]
    rows = []
    for i in range(DATA_START, ws.max_row + 1):
        first = ws.cell(i, 1).value
        if isinstance(first, str) and first.startswith("▼"):
            continue
        rec, empty = {}, True
        for j, h in enumerate(headers, start=1):
            if not h:
                continue
            v = ws.cell(i, j).value
            v = "" if v is None else str(v).strip()
            rec[h] = v
            if v:
                empty = False
        if not empty:
            rows.append(rec)
    return rows
